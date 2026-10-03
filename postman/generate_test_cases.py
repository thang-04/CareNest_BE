"""
Builds the CareNest test cases in two formats from one list:
  - CareNest.postman_collection.json  (import into Postman, run the whole collection with Runner)
  - CareNest_TestCases.xlsx           (test case document)

Run: python postman/generate_test_cases.py
"""
import json
import os
import struct
import zlib

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))

# Every run uses new emails / phone numbers, so the collection can be run many times on the same DB.
RUN_SETUP = [
    "const runId = String(Date.now()).slice(-8);",
    "const set = (k, v) => pm.collectionVariables.set(k, v);",
    "set('runId', runId);",
    "set('viceEmail', 'pho' + runId + '@gmail.com');",
    "set('vice2Email', 'pho2_' + runId + '@gmail.com');",
    "set('teacherEmail', 'gv' + runId + '@gmail.com');",
    "set('staffEmail', 'nv' + runId + '@gmail.com');",
    "set('teacherPhone', '07' + runId);",
    "set('parentPhone', '09' + runId);",
    "set('parentPhoneIntl', '+849' + runId);",
    "set('otherParentPhone', '03' + runId);",
]

CASES = []


def tc(module, name, method, path, token=None, body=None, form=None, urlencoded=None, status=200, code=0,
       checks=(), save=None, pre=None, pre_note="", data_note=None, expect_note=""):
    CASES.append(dict(module=module, name=name, method=method, path=path, token=token, body=body, form=form,
                      urlencoded=urlencoded, status=status, code=code, checks=list(checks), save=save or {},
                      pre=pre, pre_note=pre_note, data_note=data_note, expect_note=expect_note))


A, M, P, IMG, CH, PK, LK, OUT = ("Đăng nhập & token", "Quản lý tài khoản", "Mật khẩu", "Ảnh của tôi",
                                 "Trẻ & ảnh trẻ", "Người đón trẻ", "Khoá tài khoản", "Đăng xuất")
LOGIN = "/api/auth/login"
NEW_PASS = "parent123"

# ---------------- Đăng nhập & token ----------------
tc(A, "Đăng nhập sai mật khẩu", "POST", LOGIN, body={"identifier": "{{principalEmail}}", "password": "sai-mat-khau"},
   status=401, code=1101, pre=RUN_SETUP, expect_note="Tài khoản hoặc mật khẩu không chính xác")
tc(A, "Đăng nhập tài khoản không tồn tại", "POST", LOGIN, body={"identifier": "khongco@gmail.com", "password": "abc12345"},
   status=401, code=1101, expect_note="Không tiết lộ tài khoản có tồn tại hay không")
tc(A, "Đăng nhập thiếu dữ liệu", "POST", LOGIN, body={}, status=400, code=1000,
   checks=[("Báo lỗi trường identifier", "pm.expect(d.identifier).to.be.a('string')")])
tc(A, "Hiệu trưởng đăng nhập (mobile)", "POST", LOGIN, body={"identifier": "{{principalEmail}}", "password": "{{principalPassword}}"},
   checks=[("Có accessToken và refreshToken", "pm.expect(d.accessToken).to.be.a('string'); pm.expect(d.refreshToken).to.be.a('string')"),
           ("tokenType = Bearer", "pm.expect(d.tokenType).to.eql('Bearer')")],
   save={"principalToken": "d.accessToken", "principalRefresh": "d.refreshToken"},
   expect_note="Header X-Client-Type: mobile nên token nằm trong body")
tc(A, "Gọi API khi chưa đăng nhập", "GET", "/api/users/me", status=401, code=1100)
tc(A, "Gọi API với token giả", "GET", "/api/users/me", token="!abc.def.ghi", status=401, code=1104)
tc(A, "Xem thông tin tài khoản Hiệu trưởng", "GET", "/api/users/me", token="principalToken",
   checks=[("Có role PRINCIPAL", "pm.expect(d.roles).to.include('PRINCIPAL')"),
           ("Trạng thái ACTIVE", "pm.expect(d.status).to.eql('ACTIVE')")],
   save={"principalId": "d.id"})
tc(A, "Làm mới token (refresh)", "POST", "/api/auth/refresh-token", body={"refreshToken": "{{principalRefresh}}"},
   pre=["pm.collectionVariables.set('oldPrincipalRefresh', pm.collectionVariables.get('principalRefresh'));"],
   checks=[("Refresh token mới khác token cũ", "pm.expect(d.refreshToken).to.not.eql(pm.collectionVariables.get('oldPrincipalRefresh'))")],
   save={"principalToken": "d.accessToken", "principalRefresh": "d.refreshToken"})
tc(A, "Dùng lại refresh token cũ (đã bị thu hồi)", "POST", "/api/auth/refresh-token",
   body={"refreshToken": "{{oldPrincipalRefresh}}"}, status=401, code=1104)
tc(A, "Dùng refresh token thay access token", "GET", "/api/users/me", token="principalRefresh", status=401)

# ---------------- Quản lý tài khoản ----------------
tc(M, "Hiệu trưởng tạo Phó hiệu trưởng", "POST", "/api/management/staff", token="principalToken",
   body={"email": "{{viceEmail}}", "fullName": "Cô Mai", "password": "vice1234", "role": "VICE_PRINCIPAL"},
   status=201, checks=[("Role VICE_PRINCIPAL", "pm.expect(d.roles).to.eql(['VICE_PRINCIPAL'])"),
                       ("Bắt đổi mật khẩu lần đầu", "pm.expect(d.mustChangePassword).to.be.true")],
   save={"viceId": "d.id"})
tc(M, "Phó hiệu trưởng đăng nhập", "POST", LOGIN, body={"identifier": "{{viceEmail}}", "password": "vice1234"},
   checks=[("mustChangePassword = true", "pm.expect(d.mustChangePassword).to.be.true")],
   save={"viceToken": "d.accessToken"})
tc(M, "Phó HT tạo Giáo viên (email viết hoa)", "POST", "/api/management/staff", token="viceToken",
   body={"email": "GV{{runId}}@Gmail.com", "fullName": "Cô Lan", "phoneNumber": "{{teacherPhone}}",
         "password": "teacher123", "role": "TEACHER"},
   status=201, checks=[("Email được chuẩn hoá chữ thường", "pm.expect(d.email).to.eql(pm.collectionVariables.get('teacherEmail'))")],
   save={"teacherId": "d.id"})
tc(M, "Phó HT tạo Nhân viên", "POST", "/api/management/staff", token="viceToken",
   body={"email": "{{staffEmail}}", "fullName": "Chú Bảo", "password": "staff1234", "role": "STAFF"},
   status=201, checks=[("Role STAFF", "pm.expect(d.roles).to.eql(['STAFF'])")], save={"staffId": "d.id"})
tc(M, "Phó HT tạo Phó HT khác (không được phép)", "POST", "/api/management/staff", token="viceToken",
   body={"email": "x{{runId}}@gmail.com", "fullName": "X", "password": "vice1234", "role": "VICE_PRINCIPAL"},
   status=403, code=1103)
tc(M, "Tạo tài khoản Hiệu trưởng (không được phép)", "POST", "/api/management/staff", token="principalToken",
   body={"email": "y{{runId}}@gmail.com", "fullName": "Y", "password": "abc12345", "role": "PRINCIPAL"},
   status=400, code=1207)
tc(M, "Tạo tài khoản với email đã tồn tại", "POST", "/api/management/staff", token="principalToken",
   body={"email": "{{teacherEmail}}", "fullName": "Trùng", "password": "abc12345", "role": "TEACHER"},
   status=409, code=1200)
tc(M, "Tạo tài khoản với SĐT đã tồn tại", "POST", "/api/management/staff", token="principalToken",
   body={"email": "z{{runId}}@gmail.com", "fullName": "Trùng", "phoneNumber": "{{teacherPhone}}",
         "password": "abc12345", "role": "TEACHER"},
   status=409, code=1201)
tc(M, "Tạo tài khoản mật khẩu yếu", "POST", "/api/management/staff", token="principalToken",
   body={"email": "w{{runId}}@gmail.com", "fullName": "W", "password": "123", "role": "TEACHER"},
   status=400, code=1000, checks=[("Báo lỗi trường password", "pm.expect(d.password).to.be.a('string')")])
tc(M, "Tạo tài khoản email sai định dạng", "POST", "/api/management/staff", token="principalToken",
   body={"email": "khong-phai-email", "fullName": "W", "password": "abc12345", "role": "TEACHER"},
   status=400, code=1000, checks=[("Báo lỗi trường email", "pm.expect(d.email).to.be.a('string')")])
tc(M, "Tạo tài khoản thiếu role", "POST", "/api/management/staff", token="principalToken",
   body={"email": "v{{runId}}@gmail.com", "fullName": "V", "password": "abc12345"},
   status=400, code=1000, checks=[("Báo lỗi trường role", "pm.expect(d.role).to.be.a('string')")])
tc(M, "Tạo tài khoản role không tồn tại", "POST", "/api/management/staff", token="principalToken",
   body={"email": "u{{runId}}@gmail.com", "fullName": "U", "password": "abc12345", "role": "ABC"},
   status=400, code=1001)
tc(M, "Hiệu trưởng tạo Phó HT thứ 2", "POST", "/api/management/staff", token="principalToken",
   body={"email": "{{vice2Email}}", "fullName": "Thầy Nam", "password": "vice1234", "role": "VICE_PRINCIPAL"},
   status=201, save={"vice2Id": "d.id"})
tc(M, "Giáo viên đăng nhập bằng email", "POST", LOGIN, body={"identifier": "{{teacherEmail}}", "password": "teacher123"},
   save={"teacherToken": "d.accessToken", "teacherRefresh": "d.refreshToken"})
tc(M, "Giáo viên gọi API quản lý (không được phép)", "POST", "/api/management/staff", token="teacherToken",
   body={"email": "t{{runId}}@gmail.com", "fullName": "T", "password": "abc12345", "role": "TEACHER"},
   status=403, code=1103)
tc(M, "Chỉ định tổ trưởng", "PUT", "/api/management/teachers/{{teacherId}}/head-teacher", token="viceToken",
   checks=[("Có cả TEACHER và HEAD_TEACHER", "pm.expect(d.roles).to.include.members(['TEACHER', 'HEAD_TEACHER'])")])
tc(M, "Chỉ định tổ trưởng cho Nhân viên", "PUT", "/api/management/teachers/{{staffId}}/head-teacher", token="viceToken",
   status=404, code=1208)
tc(M, "Chỉ định tổ trưởng cho id không tồn tại", "PUT", "/api/management/teachers/999999/head-teacher", token="viceToken",
   status=404, code=1208)
tc(M, "Bỏ chức tổ trưởng", "DELETE", "/api/management/teachers/{{teacherId}}/head-teacher", token="viceToken",
   checks=[("Không còn HEAD_TEACHER", "pm.expect(d.roles).to.not.include('HEAD_TEACHER')")])
tc(M, "Token cũ của GV bị từ chối sau khi đổi chức vụ", "GET", "/api/users/me", token="teacherToken",
   status=401, code=1104, expect_note="Đổi role -> đăng xuất mọi nơi, phải đăng nhập lại")
tc(M, "Giáo viên đăng nhập lại", "POST", LOGIN, body={"identifier": "{{teacherEmail}}", "password": "teacher123"},
   save={"teacherToken": "d.accessToken", "teacherRefresh": "d.refreshToken"})

# ---------------- Trẻ & phụ huynh ----------------
tc(CH, "Thêm trẻ, tự tạo tài khoản phụ huynh (SĐT +84)", "POST", "/api/management/children", token="viceToken",
   body={"childName": "Bé Na", "dateOfBirth": "2022-05-01", "parentPhone": "{{parentPhoneIntl}}", "parentName": "Chị Hoa"},
   status=201, checks=[("parentAccountCreated = true", "pm.expect(d.parentAccountCreated).to.be.true"),
                       ("SĐT chuẩn hoá về 0...", "pm.expect(d.parent.phoneNumber).to.eql(pm.collectionVariables.get('parentPhone'))")],
   save={"childId": "d.id", "parentId": "d.parent.id"})
tc(CH, "Thêm trẻ thứ 2 cùng SĐT phụ huynh", "POST", "/api/management/children", token="principalToken",
   body={"childName": "Bé Bin", "parentPhone": "{{parentPhone}}", "parentName": "Chị Hoa"},
   status=201, checks=[("Gắn vào phụ huynh cũ", "pm.expect(d.parentAccountCreated).to.be.false; pm.expect(d.parent.id).to.eql(Number(pm.collectionVariables.get('parentId')))")],
   save={"child2Id": "d.id"})
tc(CH, "Thêm trẻ của phụ huynh khác", "POST", "/api/management/children", token="viceToken",
   body={"childName": "Bé Tí", "parentPhone": "{{otherParentPhone}}", "parentName": "Anh Tùng"},
   status=201, save={"otherChildId": "d.id"})
tc(CH, "Thêm trẻ thiếu tên", "POST", "/api/management/children", token="viceToken",
   body={"parentPhone": "{{parentPhone}}", "parentName": "Chị Hoa"}, status=400, code=1000,
   checks=[("Báo lỗi childName", "pm.expect(d.childName).to.be.a('string')")])
tc(CH, "Thêm trẻ SĐT phụ huynh sai", "POST", "/api/management/children", token="viceToken",
   body={"childName": "Bé", "parentPhone": "12345", "parentName": "X"}, status=400, code=1000,
   checks=[("Báo lỗi parentPhone", "pm.expect(d.parentPhone).to.be.a('string')")])
tc(CH, "Thêm trẻ ngày sinh ở tương lai", "POST", "/api/management/children", token="viceToken",
   body={"childName": "Bé", "dateOfBirth": "2999-01-01", "parentPhone": "{{parentPhone}}", "parentName": "X"},
   status=400, code=1000, checks=[("Báo lỗi dateOfBirth", "pm.expect(d.dateOfBirth).to.be.a('string')")])
tc(CH, "Giáo viên thêm trẻ (không được phép)", "POST", "/api/management/children", token="teacherToken",
   body={"childName": "Bé", "parentPhone": "{{parentPhone}}", "parentName": "X"}, status=403, code=1103)
tc(CH, "Phụ huynh đăng nhập bằng SĐT, mật khẩu mặc định", "POST", LOGIN,
   body={"identifier": "{{parentPhone}}", "password": "{{parentPhone}}"},
   checks=[("Bắt đổi mật khẩu", "pm.expect(d.mustChangePassword).to.be.true")],
   save={"parentToken": "d.accessToken"})
tc(CH, "Phụ huynh gọi API giáo viên (không được phép)", "GET", "/api/teacher/children/{{childId}}/pickup-persons",
   token="parentToken", status=403, code=1103)
tc(CH, "Phụ huynh xem danh sách con", "GET", "/api/parent/children", token="parentToken",
   checks=[("Có 2 trẻ", "pm.expect(d).to.have.lengthOf(2)")])

# ---------------- Mật khẩu ----------------
tc(P, "Đổi mật khẩu sai mật khẩu cũ", "POST", "/api/auth/change-password", token="parentToken",
   body={"oldPassword": "sai12345", "newPassword": NEW_PASS, "confirmPassword": NEW_PASS}, status=400, code=1204)
tc(P, "Đổi mật khẩu xác nhận không khớp", "POST", "/api/auth/change-password", token="parentToken",
   body={"oldPassword": "{{parentPhone}}", "newPassword": NEW_PASS, "confirmPassword": "khac1234"}, status=400, code=1202)
tc(P, "Đổi mật khẩu mới chỉ có số", "POST", "/api/auth/change-password", token="parentToken",
   body={"oldPassword": "{{parentPhone}}", "newPassword": "12345678", "confirmPassword": "12345678"},
   status=400, code=1000, checks=[("Báo lỗi newPassword", "pm.expect(d.newPassword).to.be.a('string')")],
   expect_note="Mật khẩu phải 8-100 ký tự, gồm cả chữ và số")
tc(P, "Đổi mật khẩu thành công", "POST", "/api/auth/change-password", token="parentToken",
   body={"oldPassword": "{{parentPhone}}", "newPassword": NEW_PASS, "confirmPassword": NEW_PASS})
tc(P, "Đăng nhập bằng mật khẩu cũ sau khi đổi", "POST", LOGIN,
   body={"identifier": "{{parentPhone}}", "password": "{{parentPhone}}"}, status=401, code=1101)
tc(P, "Đăng nhập bằng mật khẩu mới", "POST", LOGIN, body={"identifier": "{{parentPhone}}", "password": NEW_PASS},
   checks=[("Không còn bắt đổi mật khẩu", "pm.expect(d.mustChangePassword).to.be.false")],
   save={"parentToken": "d.accessToken"})
tc(P, "Đổi mật khẩu mới trùng mật khẩu cũ", "POST", "/api/auth/change-password", token="parentToken",
   body={"oldPassword": NEW_PASS, "newPassword": NEW_PASS, "confirmPassword": NEW_PASS}, status=400, code=1205)
tc(P, "Giáo viên reset mật khẩu phụ huynh không quản lý", "POST", "/api/teacher/parents/{{parentId}}/reset-password",
   token="teacherToken", status=403, code=1206, expect_note="Trẻ do BGH thêm nên GV không quản lý (chờ entity Lớp)")
tc(P, "BGH reset mật khẩu phụ huynh về SĐT", "POST", "/api/management/parents/{{parentId}}/reset-password", token="viceToken")
tc(P, "Phụ huynh đăng nhập lại bằng mật khẩu mặc định", "POST", LOGIN,
   body={"identifier": "{{parentPhone}}", "password": "{{parentPhone}}"},
   checks=[("Bắt đổi mật khẩu lại", "pm.expect(d.mustChangePassword).to.be.true")],
   save={"parentToken": "d.accessToken", "parentRefresh": "d.refreshToken"})
tc(P, "Reset mật khẩu cho tài khoản không phải phụ huynh", "POST", "/api/management/parents/{{teacherId}}/reset-password",
   token="viceToken", status=404, code=1203)
tc(P, "Quên mật khẩu bằng SĐT (chưa hỗ trợ SMS)", "POST", "/api/auth/forgot-password",
   body={"identifier": "{{parentPhone}}"}, status=501, code=1304)

# ---------------- Ảnh của tôi ----------------
tc(IMG, "Đổi ảnh đại diện", "PUT", "/api/users/me/avatar", token="teacherToken", form=[("file", "@sample.png")],
   checks=[("Có avatarUrl", "pm.expect(d.avatarUrl).to.include('avatars/')")])
tc(IMG, "Upload file không phải ảnh", "PUT", "/api/users/me/avatar", token="teacherToken", form=[("file", "@not-image.png")],
   status=400, code=1400, expect_note="Server kiểm tra nội dung file, không tin đuôi .png")
tc(IMG, "Upload thiếu file", "POST", "/api/users/me/images", token="teacherToken", form=[("title", "Không có file")],
   status=400, code=1000)
tc(IMG, "Thêm ảnh có tiêu đề", "POST", "/api/users/me/images", token="teacherToken",
   form=[("file", "@sample.png"), ("title", "Chứng chỉ")], status=201,
   checks=[("title đúng", "pm.expect(d.title).to.eql('Chứng chỉ')"), ("Có url", "pm.expect(d.url).to.be.a('string')")],
   save={"imageId": "d.id"})
tc(IMG, "Xem danh sách ảnh của tôi", "GET", "/api/users/me/images", token="teacherToken",
   checks=[("Có 1 ảnh", "pm.expect(d).to.have.lengthOf(1)")])
tc(IMG, "BGH xem ảnh của giáo viên", "GET", "/api/management/users/{{teacherId}}/images", token="principalToken",
   checks=[("Có 1 ảnh", "pm.expect(d).to.have.lengthOf(1)")])
tc(IMG, "Giáo viên xem ảnh người khác qua API quản lý", "GET", "/api/management/users/{{viceId}}/images",
   token="teacherToken", status=403, code=1103)
tc(IMG, "Xoá ảnh của người khác", "DELETE", "/api/users/me/images/{{imageId}}", token="principalToken",
   status=404, code=1403)
tc(IMG, "Xoá ảnh của mình", "DELETE", "/api/users/me/images/{{imageId}}", token="teacherToken")
tc(IMG, "Xem tài khoản sau khi đổi avatar", "GET", "/api/users/me", token="teacherToken",
   checks=[("Có avatarUrl", "pm.expect(d.avatarUrl).to.be.a('string')")])

# ---------------- Ảnh trẻ ----------------
tc(CH, "Phụ huynh đổi ảnh con", "PUT", "/api/parent/children/{{childId}}/photo", token="parentToken",
   form=[("file", "@sample.png")], checks=[("Có photoUrl", "pm.expect(d.photoUrl).to.include('children/')")])
tc(CH, "Phụ huynh đổi ảnh trẻ không phải con mình", "PUT", "/api/parent/children/{{otherChildId}}/photo",
   token="parentToken", form=[("file", "@sample.png")], status=404, code=1500)
tc(CH, "BGH đổi ảnh trẻ", "PUT", "/api/management/children/{{child2Id}}/photo", token="principalToken",
   form=[("file", "@sample.png")], checks=[("Có photoUrl", "pm.expect(d.photoUrl).to.be.a('string')")])

# ---------------- Người đón trẻ ----------------
PICK = [("fullName", "Bà Tư"), ("relationship", "Bà ngoại"), ("phoneNumber", "0933333333")]
tc(PK, "Phụ huynh thêm người đón", "POST", "/api/parent/children/{{childId}}/pickup-persons", token="parentToken",
   form=PICK + [("photo", "@sample.png")], status=201,
   checks=[("Trạng thái PENDING", "pm.expect(d.status).to.eql('PENDING')"), ("Có photoUrl", "pm.expect(d.photoUrl).to.be.a('string')")],
   save={"pickupPersonId": "d.id"})
tc(PK, "Thêm người đón thiếu ảnh", "POST", "/api/parent/children/{{childId}}/pickup-persons", token="parentToken",
   form=PICK, status=400, code=1000)
tc(PK, "Thêm người đón SĐT sai", "POST", "/api/parent/children/{{childId}}/pickup-persons", token="parentToken",
   form=[("fullName", "Bà Tư"), ("relationship", "Bà ngoại"), ("phoneNumber", "123"), ("photo", "@sample.png")],
   status=400, code=1000, checks=[("Báo lỗi phoneNumber", "pm.expect(d.phoneNumber).to.be.a('string')")])
tc(PK, "Gửi form-urlencoded thay vì multipart", "POST", "/api/parent/children/{{childId}}/pickup-persons",
   token="parentToken", urlencoded=PICK, status=415, code=415,
   expect_note="Chưa khớp endpoint nên trả theo format GlobalExceptionHandler: code = HTTP status")
tc(PK, "Thêm người đón cho trẻ không phải con mình", "POST", "/api/parent/children/{{otherChildId}}/pickup-persons",
   token="parentToken", form=PICK + [("photo", "@sample.png")], status=404, code=1500)
tc(PK, "Phụ huynh xem danh sách người đón", "GET", "/api/parent/children/{{childId}}/pickup-persons", token="parentToken",
   checks=[("Có 1 người", "pm.expect(d).to.have.lengthOf(1)")])
tc(PK, "Giáo viên xem người đón khi chưa duyệt", "GET", "/api/teacher/children/{{childId}}/pickup-persons",
   token="teacherToken", checks=[("Chưa thấy ai", "pm.expect(d).to.have.lengthOf(0)")])
tc(PK, "BGH xem danh sách chờ duyệt", "GET", "/api/management/pickup-persons?status=PENDING", token="viceToken",
   checks=[("Có người vừa thêm", "pm.expect(d.map(p => p.id)).to.include(Number(pm.collectionVariables.get('pickupPersonId')))")])
tc(PK, "Lọc theo trạng thái không hợp lệ", "GET", "/api/management/pickup-persons?status=ABC", token="viceToken",
   status=400, code=1000)
tc(PK, "Phụ huynh tự duyệt (không được phép)", "POST", "/api/management/pickup-persons/{{pickupPersonId}}/approve",
   token="parentToken", status=403, code=1103)
tc(PK, "Giáo viên duyệt (không được phép)", "POST", "/api/management/pickup-persons/{{pickupPersonId}}/approve",
   token="teacherToken", status=403, code=1103)
tc(PK, "BGH duyệt người đón", "POST", "/api/management/pickup-persons/{{pickupPersonId}}/approve", token="viceToken",
   checks=[("APPROVED", "pm.expect(d.status).to.eql('APPROVED')"), ("Có reviewedAt", "pm.expect(d.reviewedAt).to.be.a('string')")])
tc(PK, "Giáo viên xem người đón đã duyệt", "GET", "/api/teacher/children/{{childId}}/pickup-persons", token="teacherToken",
   checks=[("Thấy 1 người", "pm.expect(d).to.have.lengthOf(1)")])
tc(PK, "Phụ huynh sửa người đón (giữ ảnh cũ)", "PUT", "/api/parent/pickup-persons/{{pickupPersonId}}", token="parentToken",
   form=[("fullName", "Bà Tư"), ("relationship", "Bà nội"), ("phoneNumber", "0933333333")],
   checks=[("Quay về PENDING", "pm.expect(d.status).to.eql('PENDING')"), ("relationship mới", "pm.expect(d.relationship).to.eql('Bà nội')")])
tc(PK, "Giáo viên không còn thấy người đón vừa sửa", "GET", "/api/teacher/children/{{childId}}/pickup-persons",
   token="teacherToken", checks=[("Không thấy ai", "pm.expect(d).to.have.lengthOf(0)")])
tc(PK, "BGH từ chối có lý do", "POST", "/api/management/pickup-persons/{{pickupPersonId}}/reject", token="viceToken",
   body={"reason": "Ảnh không rõ mặt"},
   checks=[("REJECTED", "pm.expect(d.status).to.eql('REJECTED')"), ("Có lý do", "pm.expect(d.rejectReason).to.eql('Ảnh không rõ mặt')")])
tc(PK, "Duyệt người đón không tồn tại", "POST", "/api/management/pickup-persons/999999/approve", token="viceToken",
   status=404, code=1501)
tc(PK, "Phụ huynh xoá người đón", "DELETE", "/api/parent/pickup-persons/{{pickupPersonId}}", token="parentToken")
tc(PK, "Xoá người đón đã bị xoá", "DELETE", "/api/parent/pickup-persons/{{pickupPersonId}}", token="parentToken",
   status=404, code=1501)

# ---------------- Khoá tài khoản ----------------
tc(LK, "Phó HT khoá Phó HT khác (không được phép)", "PUT", "/api/management/users/{{vice2Id}}/status", token="viceToken",
   body={"status": "LOCKED"}, status=403, code=1103)
tc(LK, "Phó HT khoá Hiệu trưởng (không được phép)", "PUT", "/api/management/users/{{principalId}}/status", token="viceToken",
   body={"status": "LOCKED"}, status=403, code=1103)
tc(LK, "Hiệu trưởng tự khoá mình", "PUT", "/api/management/users/{{principalId}}/status", token="principalToken",
   body={"status": "LOCKED"}, status=400, code=1209)
tc(LK, "Giáo viên khoá người khác (không được phép)", "PUT", "/api/management/users/{{staffId}}/status",
   token="teacherToken", body={"status": "LOCKED"}, status=403, code=1103)
tc(LK, "Trạng thái không hợp lệ", "PUT", "/api/management/users/{{teacherId}}/status", token="principalToken",
   body={"status": "XYZ"}, status=400, code=1001)
tc(LK, "Phó HT khoá giáo viên", "PUT", "/api/management/users/{{teacherId}}/status", token="viceToken",
   body={"status": "LOCKED"}, checks=[("LOCKED", "pm.expect(d.status).to.eql('LOCKED')")])
tc(LK, "Giáo viên bị khoá dùng token cũ", "GET", "/api/users/me", token="teacherToken", status=403, code=1102,
   expect_note="Bị chặn ngay, không đợi token hết hạn")
tc(LK, "Giáo viên bị khoá làm mới token", "POST", "/api/auth/refresh-token", body={"refreshToken": "{{teacherRefresh}}"},
   status=403, code=1102)
tc(LK, "Giáo viên bị khoá đăng nhập", "POST", LOGIN, body={"identifier": "{{teacherEmail}}", "password": "teacher123"},
   status=403, code=1102)
tc(LK, "Hiệu trưởng mở khoá giáo viên", "PUT", "/api/management/users/{{teacherId}}/status", token="principalToken",
   body={"status": "ACTIVE"}, checks=[("ACTIVE", "pm.expect(d.status).to.eql('ACTIVE')")])
tc(LK, "Giáo viên đăng nhập lại sau khi mở khoá", "POST", LOGIN,
   body={"identifier": "{{teacherEmail}}", "password": "teacher123"},
   save={"teacherToken": "d.accessToken", "teacherRefresh": "d.refreshToken"})
tc(LK, "Phó HT khoá phụ huynh", "PUT", "/api/management/users/{{parentId}}/status", token="viceToken",
   body={"status": "LOCKED"})
tc(LK, "Phụ huynh bị khoá gọi API", "GET", "/api/parent/children", token="parentToken", status=403, code=1102)
tc(LK, "Phó HT mở khoá phụ huynh", "PUT", "/api/management/users/{{parentId}}/status", token="viceToken",
   body={"status": "ACTIVE"})

# ---------------- Danh sách / sửa / xoá tài khoản ----------------
U = "Sửa & xoá tài khoản"
tc(U, "Danh sách tài khoản lọc role + từ khoá", "GET", "/api/management/users?role=TEACHER&keyword=gv{{runId}}",
   token="viceToken", checks=[("Tìm thấy 1 giáo viên", "pm.expect(d.totalElements).to.eql(1)")])
tc(U, "Danh sách tài khoản có phân trang", "GET", "/api/management/users?page=0&size=2", token="principalToken",
   checks=[("Tối đa 2 dòng", "pm.expect(d.items.length).to.be.at.most(2); pm.expect(d.size).to.eql(2)")])
tc(U, "Giáo viên xem danh sách tài khoản (không được phép)", "GET", "/api/management/users", token="teacherToken",
   status=403, code=1103)
tc(U, "Chi tiết tài khoản phụ huynh kèm danh sách con", "GET", "/api/management/users/{{parentId}}", token="viceToken",
   checks=[("Có 2 trẻ", "pm.expect(d.children).to.have.lengthOf(2)")])
tc(U, "Sửa email giáo viên trùng email người khác", "PUT", "/api/management/users/{{teacherId}}", token="viceToken",
   body={"fullName": "Cô Lan", "email": "{{viceEmail}}"}, status=409, code=1200)
tc(U, "Xoá email của nhân sự", "PUT", "/api/management/users/{{staffId}}", token="viceToken",
   body={"fullName": "Chú Bảo", "email": ""}, status=400, code=1211)
tc(U, "Xoá SĐT của phụ huynh", "PUT", "/api/management/users/{{parentId}}", token="viceToken",
   body={"fullName": "Chị Hoa", "phoneNumber": ""}, status=400, code=1212)
tc(U, "Phó HT sửa Phó HT khác (không được phép)", "PUT", "/api/management/users/{{vice2Id}}", token="viceToken",
   body={"fullName": "X", "email": "{{vice2Email}}"}, status=403, code=1103)
tc(U, "Sửa thông tin giáo viên (email mới, chức vụ)", "PUT", "/api/management/users/{{teacherId}}", token="viceToken",
   body={"fullName": "Cô Lan Anh", "email": "gvmoi{{runId}}@gmail.com", "phoneNumber": "{{teacherPhone}}",
         "position": "Giáo viên lớp Mầm"},
   checks=[("Email mới", "pm.expect(d.email).to.eql('gvmoi' + pm.collectionVariables.get('runId') + '@gmail.com')"),
           ("Có chức vụ", "pm.expect(d.position).to.eql('Giáo viên lớp Mầm')")],
   save={"teacherEmail": "d.email"})
tc(U, "Token cũ bị từ chối sau khi đổi email", "GET", "/api/users/me", token="teacherToken", status=401, code=1104)
tc(U, "Giáo viên đăng nhập bằng email mới", "POST", LOGIN, body={"identifier": "{{teacherEmail}}", "password": "teacher123"},
   save={"teacherToken": "d.accessToken", "teacherRefresh": "d.refreshToken"})
tc(U, "Tự sửa họ tên của mình", "PUT", "/api/users/me", token="teacherToken", body={"fullName": "Lan Anh"},
   checks=[("Tên mới", "pm.expect(d.fullName).to.eql('Lan Anh')")])
tc(U, "Phó HT nâng giáo viên lên Phó HT (không được phép)", "PUT", "/api/management/users/{{teacherId}}/role",
   token="viceToken", body={"role": "VICE_PRINCIPAL"}, status=403, code=1103)
tc(U, "Đổi role cho tài khoản chỉ là phụ huynh", "PUT", "/api/management/users/{{parentId}}/role",
   token="principalToken", body={"role": "TEACHER"}, status=400, code=1213)
tc(U, "Tự đổi role của mình", "PUT", "/api/management/users/{{principalId}}/role", token="principalToken",
   body={"role": "TEACHER"}, status=400, code=1214)
tc(U, "Đổi giáo viên thành nhân viên", "PUT", "/api/management/users/{{teacherId}}/role", token="viceToken",
   body={"role": "STAFF"}, checks=[("Chỉ còn STAFF", "pm.expect(d.roles).to.eql(['STAFF'])")])
tc(U, "Đổi lại thành giáo viên", "PUT", "/api/management/users/{{teacherId}}/role", token="viceToken",
   body={"role": "TEACHER"}, checks=[("TEACHER", "pm.expect(d.roles).to.eql(['TEACHER'])")])
tc(U, "Giáo viên đăng nhập lại sau khi đổi role", "POST", LOGIN,
   body={"identifier": "{{teacherEmail}}", "password": "teacher123"},
   save={"teacherToken": "d.accessToken", "teacherRefresh": "d.refreshToken"})
tc(U, "Đặt mật khẩu tạm quá yếu", "POST", "/api/management/users/{{staffId}}/reset-password", token="viceToken",
   body={"password": "123"}, status=400, code=1000)
tc(U, "BGH đặt mật khẩu tạm cho nhân viên", "POST", "/api/management/users/{{staffId}}/reset-password",
   token="viceToken", body={"password": "temp1234"})
tc(U, "Nhân viên đăng nhập bằng mật khẩu tạm", "POST", LOGIN,
   body={"identifier": "{{staffEmail}}", "password": "temp1234"},
   checks=[("Bắt đổi mật khẩu", "pm.expect(d.mustChangePassword).to.be.true")])
tc(U, "Tạo nhầm tài khoản", "POST", "/api/management/staff", token="viceToken",
   body={"email": "nham{{runId}}@gmail.com", "fullName": "Tạo nhầm", "password": "abc12345", "role": "STAFF"},
   status=201, save={"mistakeId": "d.id"})
tc(U, "Xoá tài khoản tạo nhầm (chưa dùng)", "DELETE", "/api/management/users/{{mistakeId}}", token="viceToken")
tc(U, "Xem tài khoản đã xoá", "GET", "/api/management/users/{{mistakeId}}", token="viceToken", status=404, code=1203)
tc(U, "Xoá tài khoản đã đăng nhập (phải khoá thay vì xoá)", "DELETE", "/api/management/users/{{teacherId}}",
   token="viceToken", status=409, code=1210)
tc(U, "Xoá phụ huynh đang có con", "DELETE", "/api/management/users/{{parentId}}", token="principalToken",
   status=409, code=1210)
tc(U, "Phó HT xoá Hiệu trưởng (không được phép)", "DELETE", "/api/management/users/{{principalId}}", token="viceToken",
   status=403, code=1103)

# ---------------- Quản lý trẻ ----------------
K = "Quản lý trẻ"
tc(K, "Danh sách trẻ tìm theo SĐT phụ huynh", "GET", "/api/management/children?keyword={{parentPhone}}",
   token="viceToken", checks=[("2 trẻ", "pm.expect(d.totalElements).to.eql(2)")])
tc(K, "Chi tiết trẻ kèm người đón", "GET", "/api/management/children/{{childId}}", token="viceToken",
   checks=[("Có thông tin trẻ", "pm.expect(d.child.id).to.eql(Number(pm.collectionVariables.get('childId')))"),
           ("Có danh sách người đón", "pm.expect(d.pickupPersons).to.be.an('array')")])
tc(K, "Sửa thông tin trẻ", "PUT", "/api/management/children/{{childId}}", token="viceToken",
   body={"childName": "Bé Na Na", "dateOfBirth": "2022-06-01"},
   checks=[("Tên mới", "pm.expect(d.fullName).to.eql('Bé Na Na')")])
tc(K, "Sửa trẻ ngày sinh ở tương lai", "PUT", "/api/management/children/{{childId}}", token="viceToken",
   body={"childName": "Bé Na", "dateOfBirth": "2999-01-01"}, status=400, code=1000)
tc(K, "Chuyển trẻ sang phụ huynh khác (nhập nhầm SĐT)", "PUT", "/api/management/children/{{child2Id}}/parent",
   token="viceToken", body={"parentPhone": "{{otherParentPhone}}", "parentName": "Anh Tùng"},
   checks=[("Gắn vào phụ huynh đã có", "pm.expect(d.parentAccountCreated).to.be.false; pm.expect(d.parent.phoneNumber).to.eql(pm.collectionVariables.get('otherParentPhone'))")],
   expect_note="Người đón của trẻ phải được duyệt lại")
tc(K, "Phụ huynh cũ chỉ còn 1 con", "GET", "/api/parent/children", token="parentToken",
   checks=[("1 trẻ", "pm.expect(d).to.have.lengthOf(1)")])
tc(K, "Xem trẻ không tồn tại", "GET", "/api/management/children/999999", token="viceToken", status=404, code=1500)

# ---------------- Đăng xuất ----------------
tc(OUT, "Giáo viên đăng xuất", "POST", "/api/auth/logout", token="teacherToken", body={"refreshToken": "{{teacherRefresh}}"})
tc(OUT, "Dùng access token sau khi đăng xuất", "GET", "/api/users/me", token="teacherToken", status=401, code=1104)
tc(OUT, "Dùng refresh token sau khi đăng xuất", "POST", "/api/auth/refresh-token",
   body={"refreshToken": "{{teacherRefresh}}"}, status=401, code=1104)


# ======================================================================
def token_header(token):
    if not token:
        return None
    value = token[1:] if token.startswith("!") else "{{" + token + "}}"
    return {"key": "Authorization", "value": "Bearer " + value}


def postman_item(i, c):
    headers = [{"key": "X-Client-Type", "value": "mobile"}]
    h = token_header(c["token"])
    if h:
        headers.append(h)
    path, _, query = c["path"].partition("?")
    url = {"raw": "{{baseUrl}}" + c["path"], "host": ["{{baseUrl}}"], "path": [p for p in path.split("/") if p]}
    if query:
        url["query"] = [{"key": k, "value": v} for k, v in (kv.split("=") for kv in query.split("&"))]
    r = {"method": c["method"], "header": headers, "url": url}
    if c["body"] is not None:
        headers.append({"key": "Content-Type", "value": "application/json"})
        r["body"] = {"mode": "raw", "raw": json.dumps(c["body"], ensure_ascii=False, indent=2)}
    elif c["form"] is not None:
        fd = []
        for k, v in c["form"]:
            if v.startswith("@"):
                fd.append({"key": k, "type": "file", "src": v[1:]})
            else:
                fd.append({"key": k, "value": v, "type": "text"})
        r["body"] = {"mode": "formdata", "formdata": fd}
    elif c["urlencoded"] is not None:
        r["body"] = {"mode": "urlencoded", "urlencoded": [{"key": k, "value": v} for k, v in c["urlencoded"]]}

    test = ["let b = {}; try { b = pm.response.json(); } catch (e) {}", "const d = b.data;",
            f"pm.test('HTTP {c['status']}', () => pm.response.to.have.status({c['status']}));"]
    if c["status"] < 300 or c["code"]:
        test.append(f"pm.test('code = {c['code']}', () => pm.expect(b.code).to.eql({c['code']}));")
    for desc, expr in c["checks"]:
        test.append(f"pm.test({json.dumps(desc, ensure_ascii=False)}, () => {{ {expr}; }});")
    for var, expr in c["save"].items():
        test.append(f"try {{ if ({expr} !== undefined) pm.collectionVariables.set('{var}', {expr}); }} catch (e) {{}}")
    events = [{"listen": "test", "script": {"type": "text/javascript", "exec": test}}]
    if c["pre"]:
        events.insert(0, {"listen": "prerequest", "script": {"type": "text/javascript", "exec": c["pre"]}})
    return {"name": f"TC{i:03d} - {c['name']}", "event": events, "request": r}


def build_collection():
    folders, by_module = [], {}
    for i, c in enumerate(CASES, 1):
        if c["module"] not in by_module:
            by_module[c["module"]] = {"name": c["module"], "item": []}
            folders.append(by_module[c["module"]])
        by_module[c["module"]]["item"].append(postman_item(i, c))
    # Runner keeps folder order, so modules that interleave must keep the global order: one flat list in run order.
    ordered = [postman_item(i, c) for i, c in enumerate(CASES, 1)]
    variables = {"baseUrl": "http://localhost:8080", "principalEmail": "admin@carenest.com",
                 "principalPassword": "Admin@123"}
    return {
        "info": {
            "name": "CareNest API - Test cases",
            "description": ("Chạy cả collection bằng Runner theo đúng thứ tự TC001 -> TC" + f"{len(CASES):03d}"
                            + ". Mỗi lần chạy tự sinh email/SĐT mới. Cần đặt Postman Settings > Working directory "
                              "= thư mục postman/ để Runner tìm thấy sample.png và not-image.png."),
            "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json",
        },
        "item": ordered,
        "variable": [{"key": k, "value": v} for k, v in variables.items()],
    }


def describe_input(c):
    if c["data_note"]:
        return c["data_note"]
    if c["body"] is not None:
        return json.dumps(c["body"], ensure_ascii=False)
    if c["form"] is not None:
        return "multipart: " + ", ".join(f"{k}={v[1:] + ' (file)' if v.startswith('@') else v}" for k, v in c["form"])
    if c["urlencoded"] is not None:
        return "form-urlencoded: " + ", ".join(f"{k}={v}" for k, v in c["urlencoded"])
    return ""


ROLE_NAMES = {"principalToken": "Hiệu trưởng", "viceToken": "Phó hiệu trưởng", "teacherToken": "Giáo viên",
              "parentToken": "Phụ huynh", "principalRefresh": "Refresh token HT", "!abc.def.ghi": "Token giả"}


def build_excel(path):
    font = "Arial"
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    wb = Workbook()

    ws = wb.active
    ws.title = "Test cases"
    headers = ["STT", "Mã TC", "Chức năng", "Tên test case", "Method", "API", "Người thực hiện (token)",
               "Dữ liệu đầu vào", "HTTP mong đợi", "Mã code mong đợi", "Kiểm tra thêm",
               "Kết quả thực tế", "Trạng thái", "Ghi chú"]
    widths = [6, 9, 18, 42, 9, 52, 20, 60, 10, 11, 45, 22, 12, 30]
    ws.append(headers)
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w
        cell = ws.cell(row=1, column=col)
        cell.font = Font(name=font, bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for i, c in enumerate(CASES, 1):
        checks = "; ".join(desc for desc, _ in c["checks"])
        if c["expect_note"]:
            checks = (checks + "; " if checks else "") + c["expect_note"]
        ws.append([i, f"TC{i:03d}", c["module"], c["name"], c["method"], c["path"],
                   ROLE_NAMES.get(c["token"], "Không token") if c["token"] else "Không token",
                   describe_input(c), c["status"], c["code"], checks, None, None, None])
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=i + 1, column=col)
            cell.font = Font(name=font, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if col in (1, 2, 5, 9, 10, 13) else "left")
            cell.border = border
            if col in (12, 13, 14):
                cell.fill = input_fill
    last = len(CASES) + 1
    dv = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Not run"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M2:M{last}")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:N{last}"

    s = wb.create_sheet("Tổng hợp")
    s.column_dimensions["A"].width = 28
    s.column_dimensions["B"].width = 12
    for col in "CDEF":
        s.column_dimensions[col].width = 12
    s["A1"] = "Tổng hợp kết quả test"
    s["A1"].font = Font(name=font, bold=True, size=14)
    s["A2"] = "Ô vàng ở sheet 'Test cases' (Kết quả thực tế, Trạng thái, Ghi chú) là ô cần điền; bảng dưới tự tính."
    s["A2"].font = Font(name=font, italic=True, size=9)
    rows = [("Chức năng", "Tổng", "Pass", "Fail", "Blocked", "Not run")]
    modules = list(dict.fromkeys(c["module"] for c in CASES))
    rng = f"'Test cases'!$C$2:$C${last}"
    st = f"'Test cases'!$M$2:$M${last}"
    for r, m in enumerate(modules, 5):
        rows.append((m, f'=COUNTIF({rng},A{r})', f'=COUNTIFS({rng},A{r},{st},"Pass")',
                     f'=COUNTIFS({rng},A{r},{st},"Fail")', f'=COUNTIFS({rng},A{r},{st},"Blocked")',
                     f'=B{r}-C{r}-D{r}-E{r}'))
    total_row = 5 + len(modules)
    rows.append(("Tổng", f"=SUM(B5:B{total_row - 1})", f"=SUM(C5:C{total_row - 1})", f"=SUM(D5:D{total_row - 1})",
                 f"=SUM(E5:E{total_row - 1})", f"=SUM(F5:F{total_row - 1})"))
    for r, row in enumerate(rows, 4):
        for col, v in enumerate(row, 1):
            cell = s.cell(row=r, column=col, value=v)
            cell.font = Font(name=font, bold=(r == 4 or r == total_row))
            cell.border = border
            if r == 4:
                cell.fill = head_fill
                cell.font = Font(name=font, bold=True, color="FFFFFF")
    s.cell(row=total_row + 1, column=1, value="Tỉ lệ Pass").font = Font(name=font, bold=True)
    pct = s.cell(row=total_row + 1, column=3, value=f"=IF(B{total_row}=0,0,C{total_row}/B{total_row})")
    pct.number_format = "0.0%"
    pct.font = Font(name=font, bold=True)

    a = wb.create_sheet("Tài khoản & dữ liệu")
    a.column_dimensions["A"].width = 26
    a.column_dimensions["B"].width = 70
    info = [
        ("Thông tin", "Giá trị"),
        ("Base URL", "http://localhost:8080"),
        ("Hiệu trưởng", "admin@carenest.com / Admin@123 (tạo sẵn khi start app)"),
        ("Phó hiệu trưởng", "pho<runId>@gmail.com / vice1234 (TC tạo)"),
        ("Giáo viên", "gv<runId>@gmail.com / teacher123 (TC tạo)"),
        ("Nhân viên", "nv<runId>@gmail.com / staff1234 (TC tạo)"),
        ("Phụ huynh", "09<runId> / mật khẩu mặc định = SĐT (TC tạo khi thêm trẻ)"),
        ("runId", "8 số cuối của thời điểm chạy, sinh ở TC001 để mỗi lần chạy không trùng dữ liệu"),
        ("File mẫu", "postman/sample.png (ảnh hợp lệ), postman/not-image.png (không phải ảnh)"),
        ("Header mobile", "X-Client-Type: mobile -> token trả trong body; không có header -> token trong cookie (web)"),
    ]
    for r, row in enumerate(info, 1):
        for col, v in enumerate(row, 1):
            cell = a.cell(row=r, column=col, value=v)
            cell.font = Font(name=font, bold=(r == 1), color="FFFFFF" if r == 1 else "000000")
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r == 1:
                cell.fill = head_fill

    codes = wb.create_sheet("Mã lỗi")
    codes.column_dimensions["A"].width = 10
    codes.column_dimensions["B"].width = 10
    codes.column_dimensions["C"].width = 70
    code_rows = [("Code", "HTTP", "Ý nghĩa"), (0, 200, "Thành công"), (1000, 400, "Dữ liệu không hợp lệ (data chứa lỗi từng trường)"),
                 (1001, 400, "Dữ liệu sai định dạng (JSON hỏng, enum sai)"), (415, 415, "Sai Content-Type, cần multipart/form-data (format GlobalExceptionHandler: {code, desc, data})"),
                 (1100, 401, "Chưa đăng nhập"), (1101, 401, "Sai tài khoản hoặc mật khẩu"), (1102, 403, "Tài khoản đã bị khoá"),
                 (1103, 403, "Không có quyền"), (1104, 401, "Token không hợp lệ / hết hạn / đã thu hồi"),
                 (1200, 409, "Email đã được sử dụng"), (1201, 409, "SĐT đã được sử dụng"), (1202, 400, "Xác nhận mật khẩu không khớp"),
                 (1203, 404, "Tài khoản không tồn tại"), (1204, 400, "Mật khẩu hiện tại sai"), (1205, 400, "Mật khẩu mới trùng mật khẩu cũ"),
                 (1206, 403, "Phụ huynh không thuộc quản lý"), (1207, 400, "Role không được phép tạo"),
                 (1208, 404, "Giáo viên không tồn tại"), (1209, 400, "Không tự khoá/mở khoá mình"),
                 (1210, 409, "Tài khoản đã có dữ liệu, không xoá được (hãy khoá)"),
                 (1211, 400, "Nhân sự phải có email"), (1212, 400, "Phụ huynh phải có SĐT"),
                 (1213, 400, "Chỉ đổi role cho tài khoản nhân sự"), (1214, 400, "Không thao tác trên chính mình"),
                 (1304, 501, "Chưa hỗ trợ OTP qua SMS"), (1400, 400, "File không phải ảnh JPG/PNG/WebP"),
                 (1401, 413, "Ảnh quá 5MB"), (1402, 503, "Lỗi lưu ảnh"), (1403, 404, "Ảnh không tồn tại"),
                 (1500, 404, "Trẻ không tồn tại / không phải con mình"), (1501, 404, "Người đón không tồn tại")]
    for r, row in enumerate(code_rows, 1):
        for col, v in enumerate(row, 1):
            cell = codes.cell(row=r, column=col, value=v)
            cell.font = Font(name=font, bold=(r == 1), color="FFFFFF" if r == 1 else "000000")
            cell.border = border
            if r == 1:
                cell.fill = head_fill
    wb.save(path)


def write_sample_files():
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b""))
    with open(os.path.join(HERE, "sample.png"), "wb") as f:
        f.write(png)
    with open(os.path.join(HERE, "not-image.png"), "wb") as f:
        f.write(b"this is not an image")


if __name__ == "__main__":
    write_sample_files()
    with open(os.path.join(HERE, "CareNest.postman_collection.json"), "w", encoding="utf-8") as f:
        json.dump(build_collection(), f, ensure_ascii=False, indent=2)
    build_excel(os.path.join(HERE, "CareNest_TestCases.xlsx"))
    print(f"{len(CASES)} test cases written")
